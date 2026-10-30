"""Shared report plan drives DOCX, PPTX, speech and a local evidence browser."""
import csv, html, json, math, os, re, shutil, subprocess, sys, zipfile
from pathlib import Path
from prs_core import compose, save, digest, canonical, now, metric_text
from prs_workflow import parse_outputs,workflow,OUTPUT_LABELS

ROOT=Path(__file__).resolve().parents[1]
def fact_refs(a,ids):
    fm={f['id']:f for f in a['facts']};sm={s['id']:s for s in a['sources']};result=[]
    for fid in ids:
        for e in fm[fid]['evidence']:
            s=sm[e['source_id']]
            result.append({'fact_id':fid,'title':s['title'],'source_id':s['id'],'version':s['version'],'unit':e['unit'],'quote':e['quote'],'locator':s['locator']})
    for c in a['conflicts']:
        if c.get('resolution') and set(c['fact_ids'])&set(ids):
            for e in c['resolution']['evidence']:
                s=sm[e['source_id']];result.append({'fact_id':c['resolution']['id'],'title':s['title'],'source_id':s['id'],'version':s['version'],'unit':e['unit'],'quote':e['quote'],'locator':s['locator']})
    return result
def reference_text(a,ids):
    result=[]
    for r in fact_refs(a,ids):
        text=f"{r['fact_id']}  {r['title']}  {r['unit']}  版本{r['version'][:8]}"
        if text not in result:result.append(text)
    return '\n'.join(result)
def label_status(f):
    if f.get('attribution_basis')=='self_report':return '本人自述，待交叉核实'
    return '方案内容，完成待核实' if f.get('completion_basis')=='design_only' else {'done':'文档记录已完成' if f.get('completion_basis')=='document_checkbox' else '已完成','in_progress':'进行中','planned':'计划'}[f['status']]

def make_deck(a,r):
    b=a['brief'];fm={f['id']:f for f in a['facts']};slides=[]
    def add(title,kind='text',**kw):
        if kind=='text':
            parts=[];current=[];used=0
            for para in kw.get('body',[]):
                height=max(70,max(1,math.ceil(len(para)/39))*43+14)+22
                if current and used+height>420:parts.append(current);current=[];used=0
                current.append(para);used+=height
            if current:parts.append(current)
            for n,body in enumerate(parts or [[]]):slides.append({'title':title+(' 续页' if n else ''),'kind':kind,**{**kw,'body':body}})
        elif kind=='table' and len(kw['values'])>6:
            values=kw['values']
            for start in range(1,len(values),5):
                slides.append({'title':title+(' 续页' if start>1 else ''),'kind':kind,**{**kw,'values':[values[0]]+values[start:start+5]}})
        else:slides.append({'title':title,'kind':kind,**kw})
    evidence_mode=b['report_type']=='evidence'
    add(r['title'],'cover',subtitle=f"{b['person']['name']}  {b['person']['role']}",body=['按所选资料核查，事件日期待补' if evidence_mode else b['period']['start']+' 至 '+b['period']['end'],b.get('disclosure','')],fact_ids=[])
    values=[['项目','本人工作','当前状态']]
    for p in a['projects']:
        own='；'.join(fm[i]['statement'] for i in p['self_facts'])
        values.append([p['name'],own or '待补充个人贡献证据','已完成与待核实分列' if evidence_mode else ('存在未完成事项' if p['ongoing_facts'] else '已有阶段交付')])
    add('阶段工作概览','table',values=values,widths=[260,650,230],fact_ids=[f['id'] for f in a['facts'] if f['kind'] in {'action','outcome'}],note='先说明个人负责的部分，再介绍团队交付。项目状态不等同于个人绩效等级。')
    for p in a['projects']:
        section=next(s for s in r['sections'] if s['id']==p['id'])
        metric_paragraphs={metric_text(m) for m in a['metrics'] if m['group'] in p['metric_groups']}
        paragraphs=[x for x in section['paragraphs'] if x not in metric_paragraphs]
        add(p['name'],'text',body=paragraphs,fact_ids=section['fact_ids'],note='清楚区分本人工作和团队成果。数据不足的事项保留待确认，不补写百分比。')
        for m in a['metrics']:
            if m['group'] in p['metric_groups'] and m['status']=='usable':
                add(m['group'][0],'chart',categories=['基线','观察结果'],values=[m['baseline'],m['result']],unit=m['group'][1],change_pct=m['change_pct'],population=m['group'][2],baseline_window=m['baseline_window'],result_window=m['result_window'],body=[metric_text(m),f"基线周期 {m['baseline_window']}\n结果周期 {m['result_window']}"],fact_ids=m['selected_fact_ids'],note='这是同一统计口径的前后变化，不将团队指标独立归因给个人，也不以试点样本替代全量样本。')
    cv=[['指标','出现的版本','处理结果']]
    for c in a['conflicts']:
        chosen=fm[c['selected_fact_id']]['metric']['value'] if c['selected_fact_id'] else None
        cv.append([c['group'][0],' / '.join(f'{v:g}{c["group"][1]}' for v in c['values']),f'确认采用{chosen:g}{c["group"][1]}' if chosen is not None else '尚未核定，量化成果暂不采用'])
    if len(cv)>1:add('数据冲突与采用口径','table',values=cv,widths=[320,320,500],fact_ids=list(dict.fromkeys(i for c in a['conflicts'] for i in c['fact_ids'])),note='有确认记录才解除冲突。不能因为文件更晚或数字更好看就自动采用。')
    else:add('数据口径说明','text',body=['本次没有可计算的业务成效指标。版本规则变化、配置时限和JSON示例值不作为业绩指标。' if evidence_mode else '本次资料未发现同口径指标冲突。后续继续保留指标范围和来源记录。'],fact_ids=[])
    rubric=[['考核维度','对应工作']]
    for cr in b['rubric']:rubric.append([cr['name'],'、'.join(p['name'] for p in a['projects'] if cr['id'] in p.get('rubric',[])) or '待补充'])
    add('考核维度对应工作','table',values=rubric,widths=[300,840],fact_ids=[],note='按公司提供的考核维度组织工作材料。本示例不代替管理者打分。')
    add('工作复盘','text',body=next(s['paragraphs'] for s in r['sections'] if s['id']=='limitations'),fact_ids=next(s['fact_ids'] for s in r['sections'] if s['id']=='limitations'))
    plan=next(s for s in r['sections'] if s['id']=='plan')
    add(plan['title'],'text',body=plan['paragraphs'],fact_ids=plan['fact_ids'],note='计划来自所提供的工作记录，尚未完成的任务不作为既有业绩。')
    refs=[['成果记录','资料与定位']]
    for p in a['projects']:
        for fid in p['self_facts'][:1]+p['team_facts'][:1]:
            rf=fact_refs(a,[fid])[0];refs.append([fid,rf['title']+' '+rf['unit']+' '+rf['version'][:8]])
    add('关键成果证据索引','table',values=refs,widths=[240,900],fact_ids=[f['id'] for f in a['facts']],note='完整引文及原始资料快照保存在配套证据浏览页，所有页面备注包含来源定位。')
    if b.get('presentation_outline'):
        slides=slides[:1]
        for item in b['presentation_outline']:
            if item.get('kind','text') not in {'text','table'}:raise ValueError('自定义提纲仅支持文字或表格；指标图由核定指标生成')
            if not item.get('title') or not item.get('fact_ids') or any(fid not in fm for fid in item['fact_ids']):raise ValueError('自定义提纲需要标题和有效事实编号')
            fields={k:v for k,v in item.items() if k in {'body','values','widths','fact_ids','note'}}
            add(item['title'],item.get('kind','text'),**fields)
    for n,s in enumerate(slides,1):
        s['number']=n;s['sources']=reference_text(a,s.get('fact_ids',[]))
        s['speaker_notes']=s.get('note','')+'\n\n'+'\n'.join(s.get('body',[]))+'\n\n资料依据\n'+s['sources']
    return {'slide_size':[1280,720],'font':'PingFang SC','cover_label':'成果证据核查' if evidence_mode else '个人工作述职','slides':slides,'input_fingerprint':a['input_fingerprint'],'disclosure':b.get('disclosure','')}

def docx(a,r,out):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name];st.font.name='Noto Sans CJK SC';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');st.font.color.rgb=RGBColor(0,0,0)
    # Remove theme font overrides and decorative borders inherited from the DOCX default template.
    for node in d.styles.element.iter(qn('w:rFonts')):
        for attr in list(node.attrib):
            if 'theme' in attr.lower():del node.attrib[attr]
        for attr in ['ascii','hAnsi','eastAsia','cs']:node.set(qn('w:'+attr),'Noto Sans CJK SC')
    for node in list(d.styles.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
    for name in ['Title','Subtitle']:
        for node in list(d.styles[name]._element.iter(qn('w:numPr'))):node.getparent().remove(node)
    normal=d.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.32;normal.paragraph_format.space_after=Pt(9)
    d.styles['Title'].font.size=Pt(27);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13)
    header=sec.header.paragraphs[0];header.text=r['person']['name']+'  '+r['title'];header.style=d.styles['Normal'];header.runs[0].font.size=Pt(9)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    def p(text,style=None):return d.add_paragraph(text,style)
    def refs(ids):
        groups={}
        for ref in fact_refs(a,ids):
            key=(ref['fact_id'],ref['title'],ref['version'][:8]);groups.setdefault(key,[]).append(ref['unit'])
        text='；'.join(fid+' '+title+' '+','.join(dict.fromkeys(units))+' 版本'+version for (fid,title,version),units in groups.items())
        if text:
            q=p('依据 '+text.replace('\n','；'))
            for run in q.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string('525C66')
    def table(headers,values,widths):
        t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for n,w in enumerate(widths):t.columns[n].width=Inches(w)
        for n,x in enumerate(headers):t.rows[0].cells[n].text=x
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in values:
            for c,x in zip(t.add_row().cells,row):c.text=str(x)
        for ri,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for ci,c in enumerate(row.cells):
                c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                pr.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF2' if ri==0 else ('F6F8FA' if ri%2==0 else 'FFFFFF'));pr.append(shade)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','bottom','left','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'110');el.set(qn('w:type'),'dxa');margins.append(el)
                pr.append(margins)
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                    for run in para.runs:run.font.size=Pt(10);run.bold=ri==0
        p('');return t
    p(r['title'],'Title');p(r['person']['name']+'  '+r['person']['role'],'Subtitle')
    scope=r['sections'][0]
    for text in scope['paragraphs']:
        if text:p(text)
    p('阶段工作概览','Heading 1')
    table(['项目','个人归属','资料状态'],[[pr['name'],'已有个人工作记录' if pr['self_facts'] else '待核对','已完成与待核实分列' if r['report_type']=='evidence' else ('存在待完成任务' if pr['ongoing_facts'] else '未发现明确未完成项')] for pr in a['projects']],[2.1,2.3,2.3])
    p('本文分别陈述本人工作与团队结果。已核定指标采用确认口径，尚存冲突的数据不进入量化成果。每项核心工作附来源编号，完整引文见配套证据浏览页。')
    for project in a['projects']:
        s=next(s for s in r['sections'] if s['id']==project['id']);heading=p(s['title'],'Heading 1');heading.paragraph_format.page_break_before=True
        for text in s['paragraphs']:p(text)
        refs(s['fact_ids'])
        p('原始记录摘录','Heading 2')
        allrefs=fact_refs(a,s['fact_ids']);priority=[]
        for m in a['metrics']:
            if m['group'] in project['metric_groups'] and m['status']=='usable':priority+=project['self_facts'][:1]+m['selected_fact_ids'][-1:]
        for c in a['conflicts']:
            if set(c['fact_ids'])&set(s['fact_ids']):priority+=([c['resolution']['id']] if c.get('resolution') else c['fact_ids'])
        priority+=s['fact_ids'];selected=[];seen=set()
        for fid in priority:
            ref=next((rf for rf in allrefs if rf['fact_id']==fid),None)
            if ref and fid not in seen:selected.append(ref);seen.add(fid)
        for ref in selected[:(1 if r['report_type']=='evidence' else 3)]:
            excerpt=html.unescape(re.sub(r'<[^>]+>','',ref['quote']))
            excerpt=re.sub(r'\[([^\]]+)\]\([^\)]+\)',r'\1',excerpt)
            excerpt=re.sub(r'^[#*\s]+','',excerpt).replace('**','')
            p(excerpt);q=p(ref['fact_id']+'  '+ref['title']+'  '+ref['unit']+'（此处去除格式标记；完整引文见证据页）')
            for run in q.runs:run.font.size=Pt(9)
    for sid in ['rubric','limitations','plan']:
        s=next(s for s in r['sections'] if s['id']==sid);heading=p(s['title'],'Heading 1')
        if sid=='rubric':heading.paragraph_format.page_break_before=True
        for text in s['paragraphs']:p(text)
        refs(s['fact_ids'])
    heading=p('成果证据与待确认事项','Heading 1');heading.paragraph_format.page_break_before=True
    p(f"本次读取{a['coverage']['active_sources']}份当前资料，已引用{a['coverage']['cited_sources']}份。上述范围不代表个人全部工作。重复事实合并{len(a['duplicates'])}条，报告周期之外的事实{len(a['outside_period'])}条。")
    if r['report_type']!='evidence':
        for q in a['questions']:p(q['question'])
    groups={}
    for ref in fact_refs(a,[f['id'] for f in a['facts']]):
        key=(ref['fact_id'],ref['title'],ref['version'][:8]);groups.setdefault(key,[]).append(ref['unit'])
    table(['事实','来源','位置与版本'],[[fid,title,', '.join(dict.fromkeys(units))+'  '+version] for (fid,title,version),units in groups.items()],[.7,3.8,2.3])
    d.core_properties.title=r['title'];d.core_properties.author=r['person']['name'];d.save(out/'汇报文档.docx')

def write_html(a,r,out,outputs=None):
    E=lambda x:html.escape(str(x),quote=True);fm={f['id']:f for f in a['facts']}
    rawdir=out/'sources';rawdir.mkdir()
    for s in a['sources']:
        body=''.join(f'<section id="{E(u["unit"])}"><h2>{E(u["unit"])}</h2><pre>{E(u["text"])}</pre></section>' for u in s['units'])
        source_note='定位为本次CLI Markdown行号，附件临时签名参数已去除，未下载图片。' if s['provider']=='dingtalk' else '定位按资料格式或补充记录保留；摘要和本人自述不作为独立完成证明。'
        origin_link=('<p><a href="'+E(s['locator'])+'" target="_blank" rel="noopener">打开平台原文</a></p>') if s['locator'].startswith('https://') else ''
        (rawdir/(s['id']+'.html')).write_text(f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{E(s["title"])}</title><style>body{{font:16px/1.8 system-ui;max-width:1000px;margin:40px auto;padding:20px}}pre{{white-space:pre-wrap;font:inherit}}section:target{{background:#fff2bc}}h2{{font-size:14px;color:#536575}}a{{color:#164f73}}</style><a href="../index.html">返回成果证据</a><h1>{E(s["title"])}</h1><p>版本 {E(s["version"])}<br>定位 {E(s["locator"])}</p><p>{E(source_note)}</p>{origin_link}{body}</html>',encoding='utf-8')
    facts=[]
    for f in a['facts']:
        links=[]
        for ref in fact_refs(a,[f['id']]):links.append(f'<details><summary>{E(ref["title"])} {E(ref["unit"])}</summary><blockquote>{E(ref["quote"])}</blockquote><a href="sources/{ref["source_id"]}.html#{E(ref["unit"])}" target="_blank" rel="noopener">打开原文定位</a><small>版本 {ref["version"][:12]}</small></details>')
        facts.append(f'<article data-subject="{f["subject"]}" data-status="{f["status"]}" id="{E(f["id"])}"><div class="meta">{E(f["id"])}　{E(f.get("event_date") or "事件日期未知")}　{E({"self":"本人工作","team":"团队结果","other":"其他成员"}[f["subject"]])}　{label_status(f)}</div><h3>{E(f["statement"])}</h3>{"".join(links)}</article>')
    conflicts=[]
    for c in a['conflicts']:
        title='已按确认记录解决' if c['status']=='resolved' else '待核定，量化成果暂不采用'
        reason=c['resolution']['reason'] if c['resolution'] else '不自动选择较大值或较新文件。需要负责人确认及明细。'
        choices=''.join(f'<li><a href="#{fid}">{fid}</a>　{fm[fid]["metric"]["value"]:g}{E(c["group"][1])}</li>' for fid in c['fact_ids'])
        conflicts.append(f'<article class="conflict {c["status"]}"><div class="meta">{E(c["id"])}　{title}</div><h3>{E(c["group"][0])}</h3><ul>{choices}</ul><p>{E(reason)}</p></article>')
    report=''.join(f'<section><h3>{E(s["title"])}</h3>'+''.join(f'<p>{E(p)}</p>' for p in s['paragraphs'] if p)+''.join(f'<a class="fact-link" href="#{fid}">{fid}</a> ' for fid in s['fact_ids'])+'</section>' for s in r['sections'])
    css=(ROOT/'assets/evidence.css').read_text();js=(ROOT/'assets/evidence.js').read_text()
    selected=set(parse_outputs('all' if outputs is None else outputs))
    download_items=[('summary','工作总结.md','工作总结'),('report','汇报底稿.md','汇报底稿'),('docx','汇报文档.docx','汇报文档 DOCX'),('pptx','汇报演示.pptx','可编辑 PPTX'),('pdf','汇报文档.pdf','汇报文档 PDF'),('self-eval','绩效自评草稿.txt','自评草稿'),('speaker-notes','逐页讲述提示.txt','讲述提示')]
    downloads=''.join(f'<a href="{E(filename)}" download>{E(label)}</a>' for key,filename,label in download_items if key in selected)
    head=f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{E(r["title"])} 成果与证据</title><style>{css}</style><header><a href="#">职迹 WorkTrail</a><nav><a href="#report">汇报底稿</a><a href="#conflicts">数据冲突</a><a href="#evidence">成果证据</a></nav></header><main><section class="intro"><p class="eyebrow">{E(r["period"]["label"])} / {E(r["person"]["name"])}</p><h1>{E(r["title"])}</h1><p>{E(r["disclosure"])}</p><div class="downloads"><a href="汇报文档.docx" download>汇报文档 DOCX</a><a href="汇报演示.pptx" download>可编辑 PPTX</a><a href="汇报文档.pdf" download>汇报文档 PDF</a></div><p class="meta">{a["coverage"]["active_sources"]}份资料　{len(a["facts"])}条归并后的事实　{len(a["conflicts"])}组冲突　{len(a["duplicates"])}条重复合并</p></section><section id="conflicts"><h2>数据冲突如何处理</h2>{"".join(conflicts)}</section><section id="report"><h2>汇报底稿</h2>{report}</section><section id="evidence"><h2>成果证据账本</h2><p>展开记录可看原文引文，再打开精确定位。个人工作、团队结果和未完成任务分别保留。</p><div class="filters" role="group" aria-label="事实主体"><button data-filter="all" aria-pressed="true">全部</button><button data-filter="self" aria-pressed="false">本人工作</button><button data-filter="team" aria-pressed="false">团队结果</button><button data-filter="unfinished" aria-pressed="false">未完成</button></div><label>搜索记录 <input id="search" type="search" placeholder="项目、事实编号或原文关键词"></label><p id="count" aria-live="polite"></p><div id="facts">{"".join(facts)}</div></section><section><h2>采集与更新说明</h2><p>{E(a["coverage"]["note"])}</p><p>本页为本地只读证据浏览器。修正通过项目CLI的decision命令或由Agent录入，记录确认人和原因；重新生成后同步到各份材料。</p></section></main><footer>{E(r["disclosure"])} 文件生成不表示已提交绩效系统。</footer><script>{js}</script></html>'
    # Links follow the requested outputs; evidence-only pages never advertise missing files.
    head=re.sub(r'<div class="downloads">.*?</div>',lambda _: '<div class="downloads">'+downloads+'</div>',head,count=1)
    (out/'index.html').write_text(head,encoding='utf-8')

def write_summary(a,r,out):
    parts=['# 工作内容与成效总结',r['period']['label'],r.get('disclosure',''),'本总结分别记录个人工作、项目产出与可观察成效。待核实事项不计为已完成业绩。']
    for project in a['projects']:
        section=next(s for s in r['sections'] if s['id']==project['id'])
        parts+=['## '+project['name'],*section['paragraphs']]
        if not project['metric_groups']:parts.append('成效说明：当前没有可计算的业务测量指标；不补写效率、增长或节省人力。')
        parts.append('依据：'+('、'.join(section['fact_ids']) or '尚缺有效事实记录'))
    if a['undated_facts']:parts+=['## 时间范围待确认',f"有{len(a['undated_facts'])}条记录事件日期未知；仅证据核查模式展示，正式周期不纳入。"]
    if a['conflicts']:parts+=['## 数据口径待核对',*next(s['paragraphs'] for s in r['sections'] if s['id']=='limitations')]
    if a['questions']:parts+=['## 待补充信息',*dict.fromkeys(q['question'] for q in a['questions'])]
    ids=list(dict.fromkeys(fid for p in a['projects'] for fid in next(s for s in r['sections'] if s['id']==p['id'])['fact_ids']))
    if ids:parts+=['## 来源索引',reference_text(a,ids)]
    (out/'工作总结.md').write_text('\n\n'.join(p for p in parts if p)+'\n',encoding='utf-8')

def build_content(w,out,outputs=None):
    selected=set(parse_outputs('all' if outputs is None else outputs))
    a,r=compose(w)
    if a['invalid_facts'] or a['decision_errors']:raise ValueError('存在失效事实或过期修正，先运行analyze并修复，未覆盖旧成稿')
    if not a['sources'] or (not a['facts'] and not a['undated_facts'] and not a['outside_period']):raise ValueError('尚无可核对的来源与事实；先采集，由Agent读取原文并import-facts，再生成所选产出')
    if 'summary' in selected:write_summary(a,r,out)
    if 'docx' in selected:docx(a,r,out)
    if 'evidence' in selected:
        save(out/'analysis.json',a);save(out/'report.json',r);write_html(a,r,out,selected)
        save(out/'待确认事项.json',{'questions':a['questions'],'invalid_facts':a['invalid_facts'],'decision_errors':a['decision_errors']})
        with (out/'成果证据清单.csv').open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.writer(f);writer.writerow(['事实编号','项目','主体','状态','陈述','来源','定位','版本','引文'])
            for fact in a['facts']:
                for ref in fact_refs(a,[fact['id']]):writer.writerow([fact['id'],fact['project_id'],fact['subject'],fact['status'],fact['statement'],ref['title'],ref['unit'],ref['version'],ref['quote']])
    if 'report' in selected:
        md='# '+r['title']+'\n\n'+r['disclosure']+'\n\n'
        for s in r['sections']:md+='## '+s['title']+'\n\n'+'\n\n'.join(s['paragraphs'])+'\n\n'+reference_text(a,s['fact_ids'])+'\n\n'
        (out/'汇报底稿.md').write_text(md,encoding='utf-8')
    if 'self-eval' in selected:
        own=[f for f in a['facts'] if f['subject']=='self' and f['status']=='done' and f['kind']=='action']
        self_eval=('我在本周期承担的工作包括：'+'；'.join(f['statement'] for f in own)+'。' if own else '当前资料尚不能核定本人具体贡献，不生成第一人称业绩或自评分数。')+'\n\n'+next(s['paragraphs'][0] for s in r['sections'] if s['id']=='limitations')+'\n\n自评分数待本人依据公司标准确认。'
        (out/'绩效自评草稿.txt').write_text(self_eval,encoding='utf-8')
    if 'speaker-notes' in selected:
        speech=''
        for s in make_deck(a,r)['slides']:speech+=f'第{s["number"]}页 {s["title"]}\n'+s['speaker_notes']+'\n\n'
        (out/'逐页讲述提示.txt').write_text(speech,encoding='utf-8')
    return a,r

def runtime():
    root=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies'
    return {'node':os.environ.get('PRS_NODE') or (str(root/'node/bin/node') if (root/'node/bin/node').exists() else shutil.which('node')),'modules':os.environ.get('PRS_NODE_MODULES') or str(root/'node/node_modules'),'python':sys.executable}

def recipe_hash():
    return digest([(str(p.relative_to(ROOT)),digest(p.read_bytes())) for folder in ['scripts','assets'] for p in sorted((ROOT/folder).rglob('*')) if p.is_file() and '__pycache__' not in p.parts])

def render_document(out,qa):
    renderer=os.environ.get('PRS_DOCX_RENDERER')
    if not renderer:
        base=Path.home()/'.codex/plugins/cache/openai-primary-runtime/documents'
        found=sorted(base.glob('*/skills/documents/render_docx.py'))
        if found:renderer=str(found[-1])
    if not renderer:raise ValueError('缺少DOCX渲染器。请在Codex文档运行时中构建，或设置PRS_DOCX_RENDERER；不交付未渲染文档。')
    out=Path(out).resolve();target=Path(qa).resolve()/'docx';target.mkdir(parents=True,exist_ok=True)
    config=target/'fonts.conf';cache=target/'font-cache';cache.mkdir(exist_ok=True)
    config.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>'+html.escape(str(ROOT/'assets/fonts'))+'</dir><cachedir>'+html.escape(str(cache))+'</cachedir></fontconfig>',encoding='utf-8')
    subprocess.run([sys.executable,renderer,str(out/'汇报文档.docx'),'--output_dir',str(target),'--emit_pdf','--width','1200','--height','1600'],check=True,env={**os.environ,'FONTCONFIG_FILE':str(config),'SAL_FONTPATH':str(ROOT/'assets/fonts')})
    shutil.copy2(target/'汇报文档.pdf',out/'汇报文档.pdf')

def build(w,skip_ppt=False,outputs=None):
    w=Path(w).resolve();selection=workflow(w);selected=parse_outputs(outputs if outputs is not None else selection['outputs']);requested=list(selected)
    if skip_ppt:selected=[x for x in selected if x!='pptx']
    if not selected:raise ValueError('没有要生成的产出；请明确--outputs')
    stamp=now().replace(':','').replace('.','');run=w/'builds'/stamp;out=run/'output';out.mkdir(parents=True)
    a,r=build_content(w,out,selected)
    if 'pptx' in selected:
        rt=runtime()
        if not rt['node']:raise ValueError('PPT导出缺少Node，请运行doctor')
        save(run/'deck-plan.json',make_deck(a,r))
        env={**os.environ,'PRS_NODE_MODULES':rt['modules'],'RUNTIME_NODE_MODULES':rt['modules'],'RUNTIME_NODE':rt['node'],'PRS_PYTHON':sys.executable}
        subprocess.run([rt['node'],str(ROOT/'scripts/export_deck.mjs'),str(run/'deck-plan.json'),str(run)],env=env,check=True)
    if {'docx','pdf'}&set(selected):
        intermediate=run/'document';intermediate.mkdir();docx(a,r,intermediate);render_document(intermediate,run/'qa')
        if 'pdf' in selected:shutil.copy2(intermediate/'汇报文档.pdf',out/'汇报文档.pdf')
    manifest={'input_fingerprint':a['input_fingerprint'],'workflow_hash':digest(selection),'requested_outputs':requested,'generated_outputs':selected,'recipe_hash':recipe_hash(),'built_at':now(),'complete_ppt':'pptx' in selected,'report_type':r['report_type'],'files':{str(f.relative_to(out)):digest(f.read_bytes()) for f in out.rglob('*') if f.is_file()},'publish_ready':False}
    save(out/'manifest.json',manifest)
    temp=w/'.output-next';old=w/'.output-previous';current=w/'output'
    if temp.exists():shutil.rmtree(temp)
    shutil.copytree(out,temp)
    if old.exists():shutil.rmtree(old)
    if current.exists():current.rename(old)
    try:temp.rename(current)
    except BaseException:
        if old.exists():old.rename(current)
        raise
    if old.exists():shutil.rmtree(old)
    return {'ok':True,'output':str(current),'run':str(run),'generated_outputs':selected,'slides':len(make_deck(a,r)['slides']) if 'pptx' in selected else 0,'invalid_facts':len(a['invalid_facts']),'conflicts':[{'id':c['id'],'status':c['status']} for c in a['conflicts']]}

def verify(w):
    from prs_core import analyze,confined
    w=Path(w);a=analyze(w);m=json.loads((w/'output/manifest.json').read_text());errors=[]
    if a['input_fingerprint']!=m['input_fingerprint']:errors.append('底稿发生变化，输出已过期')
    if m.get('workflow_hash') and m['workflow_hash']!=digest(workflow(w)):errors.append('资料或产出选择发生变化，需要重新采集核对或生成')
    if m.get('recipe_hash')!=recipe_hash():errors.append('生成脚本或模板已变化，需要重建')
    for name,sha in m['files'].items():
        f=confined(w/'output',name)
        if not f.exists() or digest(f.read_bytes())!=sha:errors.append('输出文件变化或缺失 '+name)
    if a['invalid_facts'] or a['decision_errors']:errors.append('存在失效事实或修正')
    if m.get('requested_outputs') and m['requested_outputs']!=m.get('generated_outputs'):errors.append('有请求产出被跳过，不能标为完整完成')
    required={'summary':'工作总结.md','report':'汇报底稿.md','evidence':'index.html','docx':'汇报文档.docx','pdf':'汇报文档.pdf','pptx':'汇报演示.pptx','self-eval':'绩效自评草稿.txt','speaker-notes':'逐页讲述提示.txt'}
    for key in m.get('generated_outputs',[]):
        if key not in required or required[key] not in m['files']:errors.append('所选产出缺失 '+key)
    return {'ok':not errors,'errors':errors,'requested_outputs':m.get('requested_outputs'),'editable_pptx_created':m['complete_ppt'],'scope':'仅核验用户所选产出、引用版本、引文定位和文件哈希；未选PPT不代表失败。视觉与语义需另行复核。'}
