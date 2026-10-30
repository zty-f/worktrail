#!/usr/bin/env python3
"""Create an explicitly fictional half-year corpus and a reviewed extraction fixture."""
from pathlib import Path
import csv,json,sys
from prs_core import *

def create(w):
    w=Path(w).resolve()
    if (w/'brief.json').exists():raise ValueError('不覆盖已有示例；请指定新目录')
    init_workspace(w);raw=w/'materials';raw.mkdir()
    save(w/'workflow.json',{'schema_version':'1.0','source':{'mode':'local','local_folders':[str(raw)]},'outputs':['summary','report','evidence','docx','pdf','pptx','self-eval','speaker-notes']})
    brief={'schema_version':'1.0','person':{'name':'林知夏','role':'产品经理'},'report_type':'halfyear','period':{'start':'2026-01-01','end':'2026-06-30','label':'2026年上半年'},'minutes':10,'disclosure':'模拟案例。人物、项目和全部业务数据均为虚构，用于检验成果归因、冲突处理和证据追溯。','projects':[{'id':'P1','name':'售后工单改版','rubric':['R1','R2']},{'id':'P2','name':'知识库治理','rubric':['R2','R3']},{'id':'P3','name':'渠道数据看板','rubric':['R1']}],'rubric':[{'id':'R1','name':'业务结果'},{'id':'R2','name':'跨团队协作'},{'id':'R3','name':'专业能力与复用'}]}
    save(w/'brief.json',brief)
    docs={
      '01_年初目标.md':('# 年初目标与项目分工\n2026-01-08 工单改版项目总负责人为周岚。林知夏负责需求梳理和业务验收，不负责研发排期及总体交付。\n2026-01-08 上半年目标为缩短全量已结售后工单的处理时长，目标值28分钟。此处为目标，不是实绩。\n2025-12-22 林知夏完成上一年度需求盘点，此项属于上一报告周期。\n'),
      '02_一月周报.md':('# 一月工作周报\n2026-01-26 林知夏完成售后工单改版的需求梳理，交付需求说明与异常流程清单。\n2026-01-30 全量已结售后工单1月基线平均处理时长40分钟，统计包含退回重开工单。\n'),
      '03_二月周报.md':('# 二月工作周报\n2026-02-06 回顾1月26日交付：林知夏完成售后工单改版的需求梳理，交付需求说明与异常流程清单。\n2026-02-23 林知夏组织业务验收用例评审并完成问题归类。\n'),
      '04_三月上线复盘.md':('# 三月上线复盘\n2026-03-20 团队完成售后工单改版上线，项目总负责人为周岚，研发由陈砚负责。\n2026-03-20 林知夏完成需求范围内的业务验收并整理上线问题清单。\n2026-03-25 试点组已结售后工单平均处理时长18分钟，仅含5个试点坐席，不代表全量坐席。\n'),
      '05_四月治理记录.md':('# 知识库治理记录\n2026-04-10 林知夏完成知识库分类方案与审核规范，并组织试点部门评审。\n2026-04-24 团队完成知识库检索入口调整并启动试点运行。\n'),
      '06_五月周报.md':('# 五月工作周报\n2026-05-18 林知夏完成渠道数据看板的指标字典和原型评审，尚未进入研发。\n2026-05-29 林知夏完成知识库治理试点反馈归类，输出维护责任清单。\n'),
      '07_六月周报.md':('# 六月工作周报\n2026-06-26 渠道数据看板仍在等待研发排期，尚未上线，不具备上线后业务结果。\n2026-06-30 周报记载知识库本轮已审核有效文档60篇，范围为试点部门6月末有效文档。\n'),
      '08_七月计划.md':('# 下阶段计划\n2026-06-30 林知夏拟在下半年推动渠道数据看板进入研发排期，排期和上线时间需与研发负责人确认。\n2026-06-30 林知夏拟建立工单指标每月口径核对机制，并完成知识库有效文档数量核验。\n'),
      '09_指标负责人确认.md':('# 指标口径确认\n2026-06-30 数据负责人沈禾确认：全量已结售后工单4至6月平均处理时长采用25分钟，包含退回重开工单；23分钟为旧导出中过滤退回重开工单的错误口径，应停用。\n2026-06-30 知识库有效文档数量仍需核验，尚不能在48与60之间确定最终值。\n'),
      '10_六月盘点.md':('# 六月知识库盘点\n2026-06-30 盘点表记载知识库本轮已审核有效文档48篇，范围为试点部门6月末有效文档。\n'),
      '11_协作反馈.md':('# 业务协作反馈\n2026-06-22 业务负责人反馈：林知夏在工单需求澄清和验收问题反馈中能明确边界并跟进闭环。此条为协作观察，不等于全部业务结果由其独立造成。\n')
    }
    for name,body in docs.items():(raw/name).write_text(body,encoding='utf-8')
    with (raw/'12_工单指标导出.csv').open('w',encoding='utf-8-sig',newline='') as f:
        wr=csv.writer(f);wr.writerow(['日期','项目','指标','周期','群体','数值','单位','备注'])
        wr.writerow(['2026-06-28','P1','平均处理时长','2026-04-01至2026-06-30','全量已结售后工单',23,'分钟','旧导出'])
        wr.writerow(['2026-06-30','P1','平均处理时长','2026-04-01至2026-06-30','全量已结售后工单',25,'分钟','包含退回重开工单'])
    # Independent formats exercise paragraph and spreadsheet-row locations.
    from docx import Document
    d=Document();d.add_heading('业务验收说明',0);d.add_paragraph('2026-03-20 林知夏在工单改版中承担需求梳理和业务验收，周岚承担项目总体推进。');d.save(raw/'13_业务验收说明.docx')
    ingest_local(w,raw);sources=packet(w)['sources'];bytitle={s['title']:s for s in sources}
    facts=[]
    def ref(title,needle):
        s=bytitle[title];u=next(u for u in s['units'] if needle in u['text']);return {'source_id':s['id'],'version':s['version'],'unit':u['unit'],'quote':u['text']}
    def add(fid,project,kind,subject,day,status,statement,title,needle,**extra):
        facts.append({'id':fid,'project_id':project,'kind':kind,'subject':subject,'event_date':day,'status':status,'statement':statement,'evidence':[ref(title,needle)],'review':{'status':'checked','reviewer':'示例资料人工语义核对'},**extra})
    add('F01','P1','role','self','2026-01-08','done','我负责需求梳理和业务验收，周岚负责项目整体推进','01_年初目标.md','项目总负责人',role_level='contributor')
    add('F02','P1','action','self','2026-01-26','done','我完成需求说明与异常流程清单','02_一月周报.md','需求梳理',dedupe_key='P1-requirements-20260126')
    add('F03','P1','action','self','2026-01-26','done','我完成需求说明与异常流程清单','03_二月周报.md','回顾1月26日',dedupe_key='P1-requirements-20260126')
    add('F04','P1','action','self','2026-02-23','done','我组织业务验收用例评审并完成问题归类','03_二月周报.md','2026-02-23')
    add('F05','P1','outcome','team','2026-03-20','done','团队完成售后工单改版上线','04_三月上线复盘.md','团队完成')
    add('F06','P1','action','self','2026-03-20','done','我完成业务验收并整理上线问题清单','04_三月上线复盘.md','林知夏完成')
    add('F07','P1','action','self','2025-12-22','done','我完成上一年度需求盘点','01_年初目标.md','2025-12-22')
    def metric(fid,project,title,needle,key,value,stage,unit,pop,comp,window,day):
        add(fid,project,'metric','team',day,'done',key+'观测记录',title,needle,metric={'key':key,'value':value,'stage':stage,'unit':unit,'population':pop,'comparison_id':comp,'window':window,'direction':'lower' if unit=='分钟' else 'higher'})
    metric('M01','P1','02_一月周报.md','40分钟','平均处理时长',40,'baseline','分钟','全量已结售后工单','H1-ticket','2026-01-01至2026-01-31','2026-01-30')
    metric('M02','P1','12_工单指标导出.csv','数值=23','平均处理时长',23,'result','分钟','全量已结售后工单','H1-ticket','2026-04-01至2026-06-30','2026-06-28')
    metric('M03','P1','12_工单指标导出.csv','数值=25','平均处理时长',25,'result','分钟','全量已结售后工单','H1-ticket','2026-04-01至2026-06-30','2026-06-30')
    metric('M04','P1','04_三月上线复盘.md','18分钟','平均处理时长',18,'result','分钟','5个试点坐席已结工单','H1-ticket','2026-03-01至2026-03-25','2026-03-25')
    add('F08','P2','action','self','2026-04-10','done','我完成分类方案与审核规范，组织试点部门评审','05_四月治理记录.md','2026-04-10')
    add('F09','P2','outcome','team','2026-04-24','done','团队调整知识库检索入口并启动试点','05_四月治理记录.md','2026-04-24')
    add('F10','P2','action','self','2026-05-29','done','我完成试点反馈归类并输出维护责任清单','06_五月周报.md','2026-05-29')
    metric('M05','P2','07_六月周报.md','60篇','已审核有效文档数量',60,'result','篇','试点部门6月末有效文档','H1-wiki','2026-06-30','2026-06-30')
    metric('M06','P2','10_六月盘点.md','48篇','已审核有效文档数量',48,'result','篇','试点部门6月末有效文档','H1-wiki','2026-06-30','2026-06-30')
    add('F11','P3','action','self','2026-05-18','done','我完成渠道数据看板的指标字典和原型评审','06_五月周报.md','2026-05-18')
    add('F12','P3','outcome','team','2026-06-26','in_progress','看板等待研发排期，尚未上线','07_六月周报.md','2026-06-26')
    add('F13','P3','plan','self','2026-06-30','planned','下半年推动看板进入研发排期，具体时间与研发负责人确认','08_七月计划.md','推动渠道')
    add('F14','P2','plan','self','2026-06-30','planned','建立工单指标每月核对机制，并核验知识库有效文档数量','08_七月计划.md','每月口径')
    save(w/'facts.json',{'schema_version':'1.0','facts':facts});import_facts(w,w/'facts.json')
    before=analyze(w);save(w/'qa/analysis-before-resolution.json',before)
    c=next(c for c in before['conflicts'] if c['group'][0]=='平均处理时长')
    resolution={'id':'D01','type':'metric_resolution','conflict_id':c['id'],'selected_fact_id':'M03','reason':'采用模拟数据负责人确认的包含退回重开工单口径；23为已停用旧口径，原始记录保留。','reviewer':'模拟数据负责人沈禾','evidence':[ref('09_指标负责人确认.md','采用25分钟')]}
    save(w/'decisions/01_工单指标口径确认.json',resolution);add_decision(w,resolution)
    after=analyze(w);save(w/'qa/analysis-after-resolution.json',after)
    bad=json.loads(canonical(facts[0]));bad['id']='BAD01';bad['statement']='我主导工单改版并独立实现团队全部业务结果';bad['role_level']='contributor'
    save(w/'qa/rejected-attribution-example.json',{'fact':bad,'errors':validate_fact(bad,{s['id']:s for s in sources},brief)})
    return {'workspace':str(w),'sources':len(sources),'raw_facts':len(facts),'after_dedup':len(after['facts']),'conflicts':[c['status'] for c in after['conflicts']]}

if __name__=='__main__':print(json.dumps(create(Path(sys.argv[1])),ensure_ascii=False,indent=2))
